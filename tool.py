#!/usr/bin/env python3
"""Tools for processing files by sample IDs listed in an Excel file."""

import argparse
import csv
import os
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from tqdm import tqdm


SAMPLE_COLUMN = "样本编号"
EXCEL_SUFFIXES = {".xlsx", ".xls"}


def cell_value_to_text(value):
    """Convert spreadsheet or CSV values to comparable text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sample_ids(excel_path, column_name=SAMPLE_COLUMN):
    """Read non-empty sample IDs from the specified Excel column."""
    excel_file = Path(excel_path)
    if not excel_file.is_file():
        raise FileNotFoundError("Excel 文件不存在: {}".format(excel_file))

    workbook = load_workbook(str(excel_file), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("Excel 文件为空: {}".format(excel_file))

        try:
            column_index = list(header).index(column_name)
        except ValueError:
            columns = ", ".join(str(name) for name in header if name is not None)
            raise ValueError(
                "Excel 中未找到列 '{}'; 当前列为: {}".format(column_name, columns)
            )

        sample_ids = []
        seen = set()
        for row in rows:
            value = row[column_index] if column_index < len(row) else None
            sample_id = cell_value_to_text(value)
            if sample_id and sample_id not in seen:
                sample_ids.append(sample_id)
                seen.add(sample_id)
    finally:
        workbook.close()

    if not sample_ids:
        raise ValueError("Excel 列 '{}' 中没有有效样本编号".format(column_name))

    return sample_ids


def build_destination_path(destination_dir, filename):
    """Return a destination path, adding a suffix if a file already exists."""
    destination = destination_dir / filename
    if not destination.exists():
        return destination

    stem = destination.stem
    suffix = destination.suffix
    counter = 1
    while True:
        candidate = destination_dir / "{}_{}{}".format(stem, counter, suffix)
        if not candidate.exists():
            return candidate
        counter += 1


def copy_matched_files(excel_path, search_path, copy_path, column_name=SAMPLE_COLUMN):
    """Copy files whose file names contain any sample ID from the Excel file."""
    sample_ids = read_sample_ids(excel_path, column_name=column_name)

    search_dir = Path(search_path)
    if not search_dir.is_dir():
        raise NotADirectoryError("查找文件路径不是有效目录: {}".format(search_dir))

    destination_dir = Path(copy_path)
    destination_dir.mkdir(parents=True, exist_ok=True)

    copied_files = []
    for current_dir, _, filenames in tqdm(
        os.walk(str(search_dir)), desc="查找文件", unit="目录"
    ):
        current_path = Path(current_dir)
        for filename in filenames:
            if any(sample_id in filename for sample_id in sample_ids):
                source = current_path / filename
                destination = build_destination_path(destination_dir, filename)
                shutil.copy2(str(source), str(destination))
                copied_files.append((source, destination))

    return copied_files


def value_matches_sample(value, sample_ids):
    """Return whether a value is contained in any Excel sample ID."""
    value_text = cell_value_to_text(value)
    if not value_text:
        return False

    return any(value_text in sample_id for sample_id in sample_ids)


def find_csv_files_by_stem_suffix(search_dir, stem_suffix):
    """Find CSV files whose extensionless file name ends with stem_suffix."""
    matched_files = []
    for current_dir, _, filenames in tqdm(
        os.walk(str(search_dir)), desc="查找CSV", unit="目录"
    ):
        current_path = Path(current_dir)
        for filename in filenames:
            file_path = current_path / filename
            if file_path.suffix.lower() == ".csv" and file_path.stem.endswith(
                stem_suffix
            ):
                matched_files.append(file_path)

    return matched_files


def filter_csv_files_by_sample_column(
    excel_path,
    search_path,
    csv_stem_suffix,
    search_column_index,
    save_path,
    column_name=SAMPLE_COLUMN,
    encoding="utf-8-sig",
):
    """Filter matched CSV files by sample IDs and save retained rows."""
    if search_column_index < 0:
        raise ValueError("查找列数必须从 0 开始，不能为负数")

    sample_ids = read_sample_ids(excel_path, column_name=column_name)

    search_dir = Path(search_path)
    if not search_dir.is_dir():
        raise NotADirectoryError("查找文件路径不是有效目录: {}".format(search_dir))

    destination_dir = Path(save_path)
    destination_dir.mkdir(parents=True, exist_ok=True)

    csv_files = find_csv_files_by_stem_suffix(search_dir, csv_stem_suffix)
    saved_files = []
    for csv_file in tqdm(csv_files, desc="处理CSV", unit="文件"):
        with csv_file.open("r", encoding=encoding, newline="") as source_file:
            rows = list(csv.reader(source_file))

        kept_rows = rows[:2]
        for row in rows[2:]:
            if search_column_index < len(row) and value_matches_sample(
                row[search_column_index], sample_ids
            ):
                kept_rows.append(row)

        destination = build_destination_path(destination_dir, csv_file.name)
        with destination.open("w", encoding=encoding, newline="") as target_file:
            writer = csv.writer(target_file)
            writer.writerows(kept_rows)
        saved_files.append(destination)

    return saved_files


def find_excel_files_by_stem_suffix(search_dir, stem_suffix):
    """Find xlsx/xls files whose extensionless file name ends with stem_suffix."""
    matched_files = []
    for current_dir, _, filenames in tqdm(
        os.walk(str(search_dir)), desc="查找Excel", unit="目录"
    ):
        current_path = Path(current_dir)
        for filename in filenames:
            file_path = current_path / filename
            if file_path.suffix.lower() in EXCEL_SUFFIXES and file_path.stem.endswith(
                stem_suffix
            ):
                matched_files.append(file_path)

    return matched_files


def read_xlsx_rows(excel_file):
    """Read all rows from the first worksheet of an xlsx file."""
    workbook = load_workbook(str(excel_file), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_xls_rows(excel_file):
    """Read all rows from the first worksheet of an xls file."""
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("读取 .xls 文件需要安装 xlrd") from exc

    workbook = xlrd.open_workbook(str(excel_file))
    worksheet = workbook.sheet_by_index(0)
    rows = []
    for row_index in range(worksheet.nrows):
        rows.append(
            [
                worksheet.cell_value(row_index, column_index)
                for column_index in range(worksheet.ncols)
            ]
        )
    return rows


def read_excel_rows(excel_file):
    suffix = excel_file.suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx_rows(excel_file)
    if suffix == ".xls":
        return read_xls_rows(excel_file)
    raise ValueError("不支持的 Excel 文件类型: {}".format(excel_file))


def write_xlsx_rows(excel_file, rows):
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(list(row))
    workbook.save(str(excel_file))


def write_xls_rows(excel_file, rows):
    try:
        import xlwt
    except ImportError as exc:
        raise ImportError("写入 .xls 文件需要安装 xlwt") from exc

    workbook = xlwt.Workbook()
    worksheet = workbook.add_sheet("Sheet1")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            worksheet.write(row_index, column_index, value)
    workbook.save(str(excel_file))


def write_excel_rows(excel_file, rows):
    suffix = excel_file.suffix.lower()
    if suffix == ".xlsx":
        write_xlsx_rows(excel_file, rows)
        return
    if suffix == ".xls":
        write_xls_rows(excel_file, rows)
        return
    raise ValueError("不支持的 Excel 文件类型: {}".format(excel_file))


def filter_excel_files_by_sample_column(
    excel_path,
    search_path,
    excel_stem_suffix,
    search_column_index,
    save_path,
    column_name=SAMPLE_COLUMN,
):
    """Filter matched Excel files by sample IDs and save retained rows."""
    if search_column_index < 0:
        raise ValueError("查找列数必须从 0 开始，不能为负数")

    sample_ids = read_sample_ids(excel_path, column_name=column_name)

    search_dir = Path(search_path)
    if not search_dir.is_dir():
        raise NotADirectoryError("查找文件路径不是有效目录: {}".format(search_dir))

    destination_dir = Path(save_path)
    destination_dir.mkdir(parents=True, exist_ok=True)

    excel_files = find_excel_files_by_stem_suffix(search_dir, excel_stem_suffix)
    saved_files = []
    for excel_file in tqdm(excel_files, desc="处理Excel", unit="文件"):
        rows = read_excel_rows(excel_file)
        kept_rows = rows[:1]
        for row in rows[1:]:
            if search_column_index < len(row) and value_matches_sample(
                row[search_column_index], sample_ids
            ):
                kept_rows.append(row)

        destination = build_destination_path(destination_dir, excel_file.name)
        write_excel_rows(destination, kept_rows)
        saved_files.append(destination)

    return saved_files


def add_excel_column_argument(parser):
    parser.add_argument(
        "--column",
        default=SAMPLE_COLUMN,
        help="Excel 中的样本编号列名，默认: {}".format(SAMPLE_COLUMN),
    )


def build_copy_parser(parser):
    parser.add_argument("excel_path", help="Excel 文件路径")
    parser.add_argument("search_path", help="需要递归查找的文件夹路径")
    parser.add_argument("copy_path", help="匹配文件复制到的文件夹路径")
    add_excel_column_argument(parser)


def parse_args():
    commands = {"copy-files", "filter-csv", "filter-excel"}
    raw_args = sys.argv[1:]

    if raw_args and raw_args[0] not in commands and not raw_args[0].startswith("-"):
        # Backward-compatible usage:
        # python3 tool.py <excel_path> <search_path> <copy_path>
        legacy_parser = argparse.ArgumentParser(
            description="根据 Excel 的样本编号列递归查找并复制匹配文件"
        )
        build_copy_parser(legacy_parser)
        args = legacy_parser.parse_args(raw_args)
        args.command = "copy-files"
        return args

    parser = argparse.ArgumentParser(description="根据 Excel 的样本编号列处理文件")
    subparsers = parser.add_subparsers(dest="command")

    copy_parser = subparsers.add_parser(
        "copy-files", help="递归查找并复制文件名包含样本编号的文件"
    )
    build_copy_parser(copy_parser)

    csv_parser = subparsers.add_parser(
        "filter-csv", help="按样本编号过滤指定后缀 CSV 并另存"
    )
    csv_parser.add_argument("excel_path", help="Excel 文件路径")
    csv_parser.add_argument("search_path", help="需要递归查找的文件夹路径")
    csv_parser.add_argument(
        "csv_stem_suffix", help="去掉扩展名后的 CSV 文件名后缀"
    )
    csv_parser.add_argument(
        "search_column_index", type=int, help="CSV 查找列数，从 0 开始"
    )
    csv_parser.add_argument("save_path", help="过滤后 CSV 另存的文件夹路径")
    add_excel_column_argument(csv_parser)
    csv_parser.add_argument(
        "--encoding",
        default="utf-8-sig",
        help="CSV 文件编码，默认: utf-8-sig",
    )

    excel_parser = subparsers.add_parser(
        "filter-excel", help="按样本编号过滤指定后缀 Excel 并另存"
    )
    excel_parser.add_argument("excel_path", help="Excel 文件路径")
    excel_parser.add_argument("search_path", help="需要递归查找的文件夹路径")
    excel_parser.add_argument(
        "excel_stem_suffix", help="去掉扩展名后的 xlsx/xls 文件名后缀"
    )
    excel_parser.add_argument(
        "search_column_index", type=int, help="Excel 查找列数，从 0 开始"
    )
    excel_parser.add_argument("save_path", help="过滤后 Excel 另存的文件夹路径")
    add_excel_column_argument(excel_parser)

    args = parser.parse_args(raw_args)
    if args.command is None:
        parser.error("需要指定子命令: copy-files、filter-csv 或 filter-excel")
    return args


def main():
    args = parse_args()
    if args.command == "filter-csv":
        saved_files = filter_csv_files_by_sample_column(
            args.excel_path,
            args.search_path,
            args.csv_stem_suffix,
            args.search_column_index,
            args.save_path,
            column_name=args.column,
            encoding=args.encoding,
        )
        print("完成，已另存 {} 个 CSV 文件到 {}".format(len(saved_files), args.save_path))
    elif args.command == "filter-excel":
        saved_files = filter_excel_files_by_sample_column(
            args.excel_path,
            args.search_path,
            args.excel_stem_suffix,
            args.search_column_index,
            args.save_path,
            column_name=args.column,
        )
        print("完成，已另存 {} 个 Excel 文件到 {}".format(len(saved_files), args.save_path))
    else:
        copied_files = copy_matched_files(
            args.excel_path,
            args.search_path,
            args.copy_path,
            column_name=args.column,
        )
        print("完成，已复制 {} 个文件到 {}".format(len(copied_files), args.copy_path))


if __name__ == "__main__":
    main()
